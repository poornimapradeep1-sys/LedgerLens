"""
report_export.py

Turns a batch's verdict list (the same list produced by
verifier.verify_all() / pipeline.run_batch()) into a shareable report.

Deliberately takes the SAME data structure the UI and the database use —
no separate logic path for "what a report contains" versus "what the
dashboard shows." One canonical list of verdict dicts feeds all three.
"""

from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from fpdf import FPDF

COLUMNS = [
    ("serial_number", "Serial No."),
    ("status", "Status"),
    ("job_card_client", "Client"),
    ("job_card_service_total", "Job Card Total"),
    ("account_book_total", "Account Book Total"),
    ("account_book_line_count", "Ledger Lines"),
    ("reason", "Reason"),
]

STATUS_COLORS = {
    "VERIFIED": "C6EFCE",
    "AMOUNT_MISMATCH": "FFC7CE",
    "MISSING_RECORD": "FFEB9C",
    "MANUAL_REVIEW": "D9D9D9",
}


def export_excel(results: list, output_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Verification Results"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, result in enumerate(results, start=2):
        status = result.get("status", "")
        fill = PatternFill(
            start_color=STATUS_COLORS.get(status, "FFFFFF"),
            end_color=STATUS_COLORS.get(status, "FFFFFF"),
            fill_type="solid",
        )
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=result.get(key))
            cell.fill = fill

    for col_idx in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 20

    wb.save(output_path)
    return output_path


def export_pdf(results: list, output_path: str) -> str:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Job Card / Account Book Verification Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    col_widths = [25, 32, 35, 35, 35, 20, 90]
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(68, 114, 196)
    pdf.set_text_color(255, 255, 255)
    for (_, label), w in zip(COLUMNS, col_widths):
        pdf.cell(w, 8, label, border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(0, 0, 0)
    fill_colors = {
        "VERIFIED": (198, 239, 206),
        "AMOUNT_MISMATCH": (255, 199, 206),
        "MISSING_RECORD": (255, 235, 156),
        "MANUAL_REVIEW": (217, 217, 217),
    }
    for result in results:
        color = fill_colors.get(result.get("status"), (255, 255, 255))
        pdf.set_fill_color(*color)
        for (key, _), w in zip(COLUMNS, col_widths):
            value = result.get(key)
            text = "" if value is None else str(value)
            if len(text) > 55:
                text = text[:52] + "..."
            pdf.cell(w, 7, text, border=1, fill=True)
        pdf.ln()

    pdf.output(output_path)
    return output_path

"""
ppt_helpers.py

Shared slide-building helpers + color tokens for the project's PPT
generator scripts (generate_ppt.py, generate_daily_ppts.py). Pulled out
of generate_ppt.py so both scripts build slides the same visual way
instead of duplicating widget code.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INDIGO = RGBColor(0x4F, 0x46, 0xE5)
NAVY = RGBColor(0x12, 0x13, 0x1C)
TEXT_DARK = RGBColor(0x11, 0x18, 0x27)
TEXT_GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF4, 0xF5, 0xF9)
GREEN = RGBColor(0x16, 0x65, 0x34)
GREEN_BG = RGBColor(0xDC, 0xFC, 0xE7)
RED = RGBColor(0x99, 0x1B, 0x1B)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)
AMBER_BG = RGBColor(0xFE, 0xF3, 0xC7)
AMBER = RGBColor(0x92, 0x40, 0x0E)
GRAY_BG = RGBColor(0xE5, 0xE7, 0xEB)
GRAY_TEXT = RGBColor(0x37, 0x41, 0x51)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def textbox(slide, left, top, width, height, text, size=18, color=TEXT_DARK,
            bold=False, align=PP_ALIGN.LEFT, font="Segoe UI", italic=False,
            anchor=None, line_spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
            run.font.name = font
    return box


def bullets(slide, left, top, width, height, items, size=15, color=TEXT_DARK,
            font="Segoe UI", space_after=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        indent = item.get("indent", 0) if isinstance(item, dict) else 0
        text = item["text"] if isinstance(item, dict) else item
        prefix = "•  " if indent == 0 else "-  "
        p.text = prefix + text
        p.space_after = Pt(space_after)
        p.level = indent
        for run in p.runs:
            run.font.size = Pt(size - indent * 1.5)
            run.font.color.rgb = color
            run.font.name = font
    return box


def rounded_box(slide, left, top, width, height, fill_color, text="", text_color=WHITE,
                 size=13, bold=True, align=PP_ALIGN.CENTER, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = text_color
        run.font.name = "Segoe UI"
    return shape


def day_header(slide, day_label, date_label, subtitle):
    set_bg(slide, LIGHT_BG)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    textbox(slide, Inches(0.5), Inches(0.12), Inches(3), Inches(0.5), day_label,
            size=20, bold=True, color=INDIGO)
    textbox(slide, Inches(0.5), Inches(0.58), Inches(8), Inches(0.4), date_label,
            size=13, color=RGBColor(0x9C, 0xA3, 0xAF))
    textbox(slide, Inches(4.3), Inches(0.15), Inches(8.7), Inches(0.85), subtitle,
            size=22, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def set_notes(slide, text):
    """Sets the speaker-notes pane for a slide (visible in PowerPoint's
    Presenter View / Notes Page view) — this is where the presentation
    "script" lives, one slide at a time, right next to the slide it's
    for instead of a separate document that can drift out of sync."""
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, page_num):
    textbox(slide, Inches(0.5), Inches(7.1), Inches(6), Inches(0.3),
            "Job Card ↔ Account Book Verification System", size=9, color=TEXT_GRAY)
    textbox(slide, Inches(12.3), Inches(7.1), Inches(0.6), Inches(0.3),
            str(page_num), size=9, color=TEXT_GRAY, align=PP_ALIGN.RIGHT)
